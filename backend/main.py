from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from fastapi.responses import FileResponse
import PyPDF2
from groq import Groq
from dotenv import load_dotenv
from pptx import Presentation
import os
import uuid

load_dotenv()

app = FastAPI()

client = Groq(api_key=os.getenv("GROQ_API_KEY"))

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

stored_text = {}

def extract_pdf_text(file):
    reader = PyPDF2.PdfReader(file)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text

def chunk_text(text, chunk_size=1200):
    words = text.split()
    chunks = []

    for i in range(0, len(words), chunk_size):
        chunks.append(" ".join(words[i:i + chunk_size]))

    return chunks
def ask_ai(prompt):
    response = client.chat.completions.create(
        model="llama-3.1-8b-instant",
        messages=[
            {"role": "system", "content": "You are a helpful AI research assistant."},
            {"role": "user", "content": prompt}
        ],
        temperature=0.3
    )

    return response.choices[0].message.content

multi_documents = {}
@app.get("/")
def home():
    return {"message": "AI Research Workspace Backend is running"}


@app.post("/upload")
async def upload_pdf(file: UploadFile = File(...)):
    text = extract_pdf_text(file.file)

    document_id = str(uuid.uuid4())
    stored_text[document_id] = text

    return {
        "document_id": document_id,
        "message": "PDF uploaded and text extracted successfully",
        "characters": len(text)
    }


@app.post("/summary")
def generate_summary(
    document_id: str = Form(...),
    summary_type: str = Form("Detailed Notes"),
    language_style: str = Form("Student-friendly"),
    output_format: str = Form("Headings + Bullet Points"),
    purpose: str = Form("Study Notes")
):
    text = stored_text.get(document_id)

    if not text:
        return {"error": "Document not found"}

    prompt = f"""
    Generate a customized document summary.

    User Preferences:
    Summary Type: {summary_type}
    Language Style: {language_style}
    Output Format: {output_format}
    Purpose: {purpose}

    Format the answer beautifully using markdown.

    Include:
    #  Customized AI Summary

    ##  Overview
    ##  Main Objective
    ##  Important Points
    ##  Methodology / Approach
    ##  Key Findings
    ##  Notes for {purpose}
    ##  Conclusion

    Rules:
    - Follow the user's selected style
    - Use clear side headings
    - Use bullet points where needed
    - Keep it useful for the selected purpose
    - Make it look like clean study notes
    - Do not say "unfortunately"
    - If document is not a research paper, still summarize professionally

    Document:
    {text[:8000]}
    """

    summary = ask_ai(prompt)
    return {"summary": summary}


@app.post("/ppt-content")
def generate_ppt_content(
    document_id: str = Form(...),
    summary_type: str = Form("Presentation Notes"),
    language_style: str = Form("Student-friendly"),
    output_format: str = Form("Slide-wise Points"),
    purpose: str = Form("Presentation")
):
    text = stored_text.get(document_id)

    if not text:
        return {"error": "Document not found"}

    prompt = f"""
    Generate presentation-ready content from this document.

    User Preferences:
    Language Style: {language_style}
    Output Format: {output_format}
    Purpose: {purpose}

    Format in markdown:

    #  Presentation Content

    ## Slide 1: Title
    - point

    ## Slide 2: Introduction
    - point
    - point

    ## Slide 3: Problem / Background
    - point
    - point

    ## Slide 4: Methodology / Approach
    - point
    - point

    ## Slide 5: Key Features / Findings
    - point
    - point

    ## Slide 6: Results / Benefits
    - point
    - point

    ## Slide 7: Applications
    - point
    - point

    ## Slide 8: Conclusion
    - point

    Rules:
    - Make it suitable for PPT slides
    - Keep points short and presentable
    - Do not generate long paragraphs

    Document:
    {text[:8000]}
    """

    ppt_content = ask_ai(prompt)
    return {"ppt_content": ppt_content}

@app.post("/upload-multiple")
async def upload_multiple_pdfs(files: list[UploadFile] = File(...)):
    project_id = str(uuid.uuid4())
    documents = []

    for file in files:
        text = extract_pdf_text(file.file)

        documents.append({
            "filename": file.filename,
            "text": text,
            "chunks": chunk_text(text)
        })

    multi_documents[project_id] = documents

    print("Project ID created:", project_id)
    print("Files uploaded:", [doc["filename"] for doc in documents])

    return {
        "project_id": project_id,
        "message": "Multiple PDFs uploaded successfully",
        "total_documents": len(documents),
        "files": [doc["filename"] for doc in documents]
    }
@app.post("/multi-ask")
def ask_across_multiple_papers(
    project_id: str = Form(...),
    question: str = Form(...)
):
    documents = multi_documents.get(project_id)

    if not documents:
        return {"error": "Project not found. Please upload multiple PDFs again."}

    all_chunks = []
    chunk_sources = []

    for doc in documents:
        for chunk in doc["chunks"]:
            if chunk.strip():
                all_chunks.append(chunk)
                chunk_sources.append(doc["filename"])

    if not all_chunks:
        return {"error": "No readable text found in uploaded PDFs."}

    vectorizer = TfidfVectorizer(stop_words="english")
    vectors = vectorizer.fit_transform(all_chunks + [question])

    question_vector = vectors[-1]
    chunk_vectors = vectors[:-1]

    similarities = cosine_similarity(question_vector, chunk_vectors).flatten()
    top_indexes = similarities.argsort()[-2:][::-1]

    context = ""

    for index in top_indexes:
        context += f"\n\nSource: {chunk_sources[index]}\nContent:\n{all_chunks[index][:1200]}"

    prompt = f"""
    You are an AI research assistant.

    Answer the question using the uploaded research papers.

    Question:
    {question}

    Context:
    {context}

    Give the answer in markdown:

    # Multi-Paper Research Analysis

    ## Direct Answer

    ## Paper-wise Insights

    ## Similarities

    ## Differences

    ## Final Conclusion

    Keep it clear, structured, and useful for students.
    """

    answer = ask_ai(prompt)

    return {
        "answer": answer,
        "sources_used": list(set([chunk_sources[i] for i in top_indexes]))
    }
@app.post("/quiz")
def generate_quiz(
    document_id: str = Form(...),
    summary_type: str = Form("Exam Notes"),
    language_style: str = Form("Student-friendly"),
    output_format: str = Form("Questions and Answers"),
    purpose: str = Form("Viva")
):
    text = stored_text.get(document_id)

    if not text:
        return {"error": "Document not found"}

    prompt = f"""
    Generate quiz and viva questions from this document.

    User Preferences:
    Language Style: {language_style}
    Purpose: {purpose}

    Format in markdown:

    #  Quiz & Viva Questions

    ##  Short Answer Questions
    1. Question?
       - Answer:

    ##  Viva Questions
    1. Question?
       - Answer:

    ##  Important Questions
    1. Question?
       - Answer:

    Rules:
    - Generate questions only from the uploaded document
    - Give clear answers
    - Keep language {language_style}
    - Make it useful for {purpose}

    Document:
    {text[:8000]}
    """

    quiz = ask_ai(prompt)
    return {"quiz": quiz}

@app.post("/ask")
def ask_question(document_id: str = Form(...), question: str = Form(...)):
    text = stored_text.get(document_id)

    if not text:
        return {"error": "Document not found"}

    prompt = f"""
    Answer the question only based on the document content.

    Document:
    {text[:8000]}

    Question:
    {question}
    """

    answer = ask_ai(prompt)
    return {"answer": answer}


@app.post("/generate-ppt")
def generate_ppt(document_id: str = Form(...)):
    text = stored_text.get(document_id)

    if not text:
        return {"error": "Document not found"}

    prompt = f"""
    Create 8 professional presentation slides from this research paper.
    For each slide give:
    Title:
    Points:
    - max 4 short bullet points

    Text:
    {text[:8000]}
    """

    content = ask_ai(prompt)

    prs = Presentation()

    title_slide_layout = prs.slide_layouts[0]
    content_slide_layout = prs.slide_layouts[1]

    dark_blue = "1F245C"
    purple = "4F46E5"

    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "AI Generated Research Presentation"
    slide.placeholders[1].text = "Created from uploaded research document"

    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
    slide.shapes.title.text_frame.paragraphs[0].font.bold = True
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(dark_blue)

    blocks = content.split("Title:")

    for block in blocks:
        if block.strip():
            lines = block.strip().split("\n")
            title = lines[0].strip()
            points = "\n".join(lines[1:]).replace("Points:", "").strip()

            slide = prs.slides.add_slide(content_slide_layout)

            fill = slide.background.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(245, 243, 255)

            slide.shapes.title.text = title
            title_frame = slide.shapes.title.text_frame
            title_frame.paragraphs[0].font.size = Pt(32)
            title_frame.paragraphs[0].font.bold = True
            title_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(dark_blue)

            body = slide.placeholders[1]
            body.text = points

            for paragraph in body.text_frame.paragraphs:
                paragraph.font.size = Pt(20)
                paragraph.font.color.rgb = RGBColor(55, 65, 81)
                paragraph.space_after = Pt(10)

            footer = slide.shapes.add_textbox(
                Inches(0.4),
                Inches(6.8),
                Inches(9),
                Inches(0.3)
            )
            footer.text = "AI Research & Presentation Workspace"
            footer.text_frame.paragraphs[0].font.size = Pt(10)
            footer.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(purple)

    os.makedirs("generated", exist_ok=True)

    file_path = f"generated/research_presentation_{document_id}.pptx"
    prs.save(file_path)

    return FileResponse(
        file_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename="AI_Research_Presentation.pptx"
    )